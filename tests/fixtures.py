"""Shared test helpers: in-memory document builders and fake providers."""

from io import BytesIO

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def pdf_bytes(*pages: str) -> bytes:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    for text in pages:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_bytes(paragraph: str = "The support window is 9 AM to 5 PM.") -> bytes:
    document = Document()
    document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Escalation"
    table.cell(0, 1).text = "Contact Mira"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


class FakeEmbedder:
    """Two-dimensional embeddings keyed on the word 'beta'; counts calls."""

    identity = "test:two-dimensional"

    def __init__(self) -> None:
        self.calls = 0

    def embed(self, texts: list[str]) -> list[list[float]]:
        self.calls += 1
        return [[0.0, 1.0] if "beta" in text else [1.0, 0.0] for text in texts]
